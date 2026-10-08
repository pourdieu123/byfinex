"""Tests de sécurité : redirections, blocages d'essais, comptes protégés, injections, en-têtes."""

from io import BytesIO

from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from . import roles, verrou
from .models import Eleve, Etablissement, remplacer_variables


def utilisateur(role, username):
    user = User.objects.create_user(username, password='Mot-de-passe-123')
    roles.definir_role(user, role)
    return user


class SecuriteTests(TestCase):
    def setUp(self):
        cache.clear()
        self.direction = utilisateur(roles.DIRECTION, 'direction')

    def test_connexion_bloquee_apres_cinq_echecs(self):
        for _i in range(5):
            self.client.post(reverse('login'), {'username': 'direction', 'password': 'faux'})
        r = self.client.post(reverse('login'), {'username': 'direction', 'password': 'Mot-de-passe-123'})
        self.assertContains(r, 'Réessayez dans')
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_code_module_bloque_meme_apres_reconnexion(self):
        verrou.definir_code(self.direction, 'scolarite', '2468')
        url = reverse('ecole:deverrouiller', args=['scolarite'])
        self.client.force_login(self.direction)
        for _i in range(5):
            self.client.post(url, {'code': '0000'})
        self.client.logout()
        self.client.force_login(self.direction)  # nouvelle session : le compteur n'est pas remis à zéro
        r = self.client.post(url, {'code': '2468', 'next': '/eleves/'})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Réessayez dans')

    @override_settings(VERROU_MODULES=False)
    def test_pas_de_redirection_vers_un_autre_site(self):
        from .facturation import enregistrer_paiement
        from .models import Facture, Paiement
        etab = Etablissement.objects.create(nom='E', emecef_actif=False)
        e = Eleve.objects.create(nom='A', prenom='B', parent='C', classe='CM2', telephone='97', montant_scolarite=1000,
                                 etablissement=etab)
        enregistrer_paiement(Paiement(eleve=e, motif='scolarite', montant=500), self.direction)
        self.client.force_login(self.direction)
        r = self.client.post(reverse('ecole:facture_normaliser', args=[Facture.objects.get().pk]),
                             {'next': 'https://site-pirate.example/'})
        self.assertRedirects(r, reverse('ecole:factures'), fetch_redirect_response=False)

    @override_settings(VERROU_MODULES=False)
    def test_direction_ne_peut_pas_modifier_un_super_administrateur(self):
        admin = User.objects.create_superuser('admin', password='Mot-de-passe-123')
        self.client.force_login(self.direction)
        url = reverse('ecole:utilisateur_modifier', args=[admin.pk])
        self.assertEqual(self.client.get(url).status_code, 403)
        self.client.post(url, {'first_name': 'x', 'last_name': 'y', 'role': roles.DIRECTION,
                               'nouveau_mot_de_passe': 'Nouveau-mot-2026'})
        admin.refresh_from_db()
        self.assertTrue(admin.check_password('Mot-de-passe-123'))

    @override_settings(VERROU_MODULES=False)
    def test_export_excel_sans_formule(self):
        etab = Etablissement.objects.create(nom='E')
        Eleve.objects.create(nom='=HYPERLINK("http://x")', prenom='B', parent='+C', classe='CM2', telephone='97',
                             montant_scolarite=1000, etablissement=etab)
        self.client.force_login(self.direction)
        classeur = load_workbook(BytesIO(self.client.get(reverse('ecole:export_eleves')).content))
        valeurs = [c.value for ligne in classeur.active.iter_rows() for c in ligne if isinstance(c.value, str)]
        self.assertFalse([v for v in valeurs if v.startswith(('=', '+'))])

    def test_variables_de_message_sans_format(self):
        self.assertEqual(remplacer_variables('{parent} {parent.__class__} {x:>99999999}', parent='Marc'),
                         'Marc {parent.__class__} {x:>99999999}')

    @override_settings(VERROU_MODULES=False)
    def test_en_tetes_de_securite(self):
        import re
        self.client.force_login(self.direction)
        r = self.client.get(reverse('ecole:eleve_nouveau'))
        nonce = re.search(r"'nonce-([^']+)'", r.headers['Content-Security-Policy']).group(1)
        self.assertIn(f'<script nonce="{nonce}">', r.content.decode())  # seuls nos scripts s'exécutent
        self.assertEqual(r.headers['X-Frame-Options'], 'DENY')
        self.assertEqual(r.headers['X-Content-Type-Options'], 'nosniff')

    def test_version_ne_revele_aucun_chiffre(self):
        self.client.force_login(self.direction)
        v = self.client.get(reverse('ecole:version')).json()['v']
        self.assertRegex(v, r'^[0-9a-f]{16}$')

    @override_settings(VERROU_MODULES=False)
    def test_parent_et_surveillant_hors_des_pages_du_personnel(self):
        parent = utilisateur(roles.PARENT, 'parent')
        surveillant = utilisateur(roles.SURVEILLANT, 'surv')
        for u in (parent, surveillant):
            self.client.force_login(u)
            for url in (reverse('ecole:eleves'), reverse('ecole:journal'), reverse('ecole:factures'),
                        reverse('ecole:utilisateurs'), reverse('ecole:messages'), reverse('ecole:grilles')):
                self.assertEqual(self.client.get(url).status_code, 403, (u.username, url))

    def test_parent_ne_voit_pas_le_tableau_de_bord(self):
        self.client.force_login(utilisateur(roles.PARENT, 'parent2'))
        self.assertEqual(self.client.get(reverse('ecole:accueil')).status_code, 403)

    @override_settings(VERROU_MODULES=False)
    def test_secretariat_ne_change_pas_le_montant_du(self):
        etab = Etablissement.objects.create(nom='E')
        e = Eleve.objects.create(nom='A', prenom='B', parent='C', classe='CM2', telephone='97', montant_scolarite=1000,
                                 etablissement=etab)
        self.client.force_login(utilisateur(roles.SECRETARIAT, 'sec'))
        self.client.post(reverse('ecole:eleve_modifier', args=[e.pk]), {
            'etablissement_nom': 'E', 'nom': 'A', 'prenom': 'B', 'classe': 'CM2', 'parent': 'C', 'telephone': '97',
            'montant_scolarite': '10'})
        self.assertEqual(Eleve.objects.get(pk=e.pk).montant_scolarite, 1000)

    @override_settings(VERROU_MODULES=False)
    def test_acces_parent_retire_quand_le_numero_change(self):
        from .models import CompteParent
        etab = Etablissement.objects.create(nom='E')
        e = Eleve.objects.create(nom='A', prenom='B', parent='C', classe='CM2', telephone='97000001',
                                 montant_scolarite=1000, etablissement=etab)
        parent = utilisateur(roles.PARENT, 'parent3')
        compte = CompteParent.objects.create(utilisateur=parent, telephone=e.telephone_international)
        compte.eleves.add(e)
        self.client.force_login(self.direction)
        self.client.post(reverse('ecole:eleve_modifier', args=[e.pk]), {
            'etablissement_nom': 'E', 'nom': 'A', 'prenom': 'B', 'classe': 'CM2', 'parent': 'C', 'telephone': '96000002',
            'montant_scolarite': '1000'})
        self.assertFalse(compte.eleves.exists())
