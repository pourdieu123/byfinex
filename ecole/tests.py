from datetime import timedelta
from io import BytesIO
from urllib.parse import unquote

from django.contrib.auth.models import Group, User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from . import roles
from .models import Eleve, Etablissement, Motif, Paiement, Relance
from .templatetags.ecole_extras import nombre_en_lettres


def creer_eleve(etablissement='École Test', **kwargs):
    data = {
        'nom': 'Dossou', 'prenom': 'Koffi', 'parent': 'Dossou Marc',
        'classe': 'CM2', 'telephone': '01 97 00 00 00', 'montant_scolarite': 60000,
        'etablissement': Etablissement.depuis_nom(etablissement),
    }
    data.update(kwargs)
    return Eleve.objects.create(**data)


def payer(eleve, montant, motif=Motif.TRANCHE1, **kwargs):
    return Paiement.objects.create(eleve=eleve, montant=montant, motif=motif, **kwargs)


def utilisateur(role, username=None):
    user = User.objects.create_user(username or role.lower(), password='Mot-de-passe-123')
    roles.definir_role(user, role)
    return user


class ModeleTests(TestCase):
    def test_statuts_et_totaux(self):
        e = creer_eleve()
        self.assertEqual(e.statut, 'aucun')
        payer(e, 20000)
        self.assertEqual(e.statut, 'partiel')
        payer(e, 40000, Motif.TRANCHE2)
        self.assertEqual((e.statut, e.reste_a_payer, e.taux_paye), ('solde', 0, 100))

    def test_inscription_hors_scolarite(self):
        e = creer_eleve()
        payer(e, 10000, Motif.INSCRIPTION)
        payer(e, 50000)
        e = Eleve.objects.avec_totaux().get(pk=e.pk)
        self.assertEqual((e.total_inscription, e.total_verse, e.reste_a_payer), (10000, 50000, 10000))

    def test_querysets_soldes_impayes(self):
        solde, partiel, aucun = creer_eleve(), creer_eleve(), creer_eleve()
        payer(solde, 60000)
        payer(partiel, 1000)
        payer(partiel, 500, Motif.INSCRIPTION)
        self.assertEqual(list(Eleve.objects.soldes()), [solde])
        self.assertCountEqual(Eleve.objects.impayes(), [partiel, aucun])
        self.assertEqual(Eleve.objects.avec_totaux().get(pk=partiel.pk).total_verse, 1000)

    def test_numero_recu_unique(self):
        e = creer_eleve()
        p1, p2 = payer(e, 100), payer(e, 200)
        self.assertRegex(p1.numero_recu, r'^R\d{4}-\d{5}$')
        self.assertNotEqual(p1.numero_recu, p2.numero_recu)

    def test_telephone_international(self):
        self.assertEqual(creer_eleve(telephone='01 97 12 34 56').telephone_international, '2290197123456')
        self.assertEqual(creer_eleve(telephone='+229 97 12 34 56').telephone_international, '22997123456')
        self.assertEqual(creer_eleve(telephone='+33 6 12 34 56 78').telephone_international, '33612345678')

    def test_nombre_en_lettres(self):
        self.assertEqual(nombre_en_lettres(71), 'soixante et onze')
        self.assertEqual(nombre_en_lettres(80), 'quatre-vingts')
        self.assertEqual(nombre_en_lettres(200000), 'deux cent mille')
        self.assertEqual(nombre_en_lettres(123456), 'cent vingt-trois mille quatre cent cinquante-six')


@override_settings(VERROU_MODULES=False, MODULES_BIENTOT=[])
class RolesTests(TestCase):
    def test_groupes_crees(self):
        self.assertEqual(sorted(roles.ROLES), sorted(g for g in roles.ROLES if roles.Group.objects.filter(name=g).exists()))

    def test_droits_par_role(self):
        e = creer_eleve()
        p = payer(e, 1000)
        pages_interdites = {
            roles.CAISSIER: [reverse('ecole:eleve_nouveau'), reverse('ecole:eleve_modifier', args=[e.pk]),
                             reverse('ecole:exports'), reverse('ecole:utilisateurs'),
                             reverse('ecole:paiement_supprimer', args=[p.pk])],
            roles.SECRETARIAT: [reverse('ecole:encaisser', args=[e.pk]), reverse('ecole:utilisateurs'),
                                reverse('ecole:eleve_archiver', args=[e.pk]), reverse('ecole:parametres')],
        }
        pages_permises = {
            roles.CAISSIER: [reverse('ecole:encaisser', args=[e.pk]), reverse('ecole:recu', args=[p.pk]),
                             reverse('ecole:relances'), reverse('ecole:journal')],
            roles.SECRETARIAT: [reverse('ecole:eleve_nouveau'), reverse('ecole:exports'), reverse('ecole:relances')],
            roles.DIRECTION: [reverse('ecole:utilisateurs'), reverse('ecole:parametres'),
                              reverse('ecole:eleve_archiver', args=[e.pk])],
        }
        for role in roles.ROLES:
            self.client.force_login(utilisateur(role))
            for url in pages_interdites.get(role, []):
                self.assertEqual(self.client.get(url).status_code, 403, f'{role} {url}')
            for url in pages_permises.get(role, []):
                self.assertEqual(self.client.get(url).status_code, 200, f'{role} {url}')

    def test_secretariat_inscrit_sans_paiements(self):
        self.client.force_login(utilisateur(roles.SECRETARIAT))
        r = self.client.get(reverse('ecole:eleve_nouveau'))
        self.assertNotIn('tranche1', r.context['form'].fields)


@override_settings(VERROU_MODULES=False, MODULES_BIENTOT=[])
class VuesTests(TestCase):
    def setUp(self):
        self.user = utilisateur(roles.DIRECTION)
        self.client.force_login(self.user)

    def test_connexion_obligatoire(self):
        self.client.logout()
        r = self.client.get(reverse('ecole:portail'))
        self.assertRedirects(r, f"{reverse('login')}?next=/")

    def test_inscription_cree_les_paiements(self):
        r = self.client.post(reverse('ecole:eleve_nouveau'), {
            'etablissement_nom': 'CEG Akpakpa', 'nom': 'Agbo', 'prenom': 'Afi', 'classe': '6ème A', 'parent': 'Agbo Rose',
            'nature_parent': 'mere', 'telephone': '96 11 22 33', 'montant_scolarite': '80000',
            'frais_inscription': '15000', 'inscription_versee': '15000', 'versement': '30000',
            'mode': 'mobile_money',
        })
        eleve = Eleve.objects.get()
        self.assertRedirects(r, reverse('ecole:eleve_detail', args=[eleve.pk]))
        self.assertEqual(eleve.paiements.count(), 2)
        self.assertTrue(all(p.mode == 'mobile_money' and p.encaisse_par == self.user for p in eleve.paiements.all()))
        self.assertEqual(eleve.reste_a_payer, 50000)

    def test_champs_obligatoires(self):
        r = self.client.post(reverse('ecole:eleve_nouveau'), {})
        for champ in ('etablissement_nom', 'nom', 'prenom', 'classe', 'parent', 'telephone', 'montant_scolarite'):
            self.assertIn(champ, r.context['form'].errors)
        self.assertFalse(Eleve.objects.exists())

    def test_tranches_superieures_a_la_scolarite(self):
        r = self.client.post(reverse('ecole:eleve_nouveau'), {
            'etablissement_nom': 'X', 'nom': 'A', 'prenom': 'B', 'classe': 'CI', 'parent': 'C', 'telephone': '1',
            'montant_scolarite': '1000', 'versement': '2000', 'mode': 'especes',
        })
        self.assertContains(r, 'dépasse le montant')
        self.assertFalse(Eleve.objects.exists())

    def test_encaisser_puis_recu(self):
        e = creer_eleve()
        r = self.client.post(reverse('ecole:encaisser', args=[e.pk]), {
            'motif': 'scolarite', 'montant': '25000', 'mode': 'especes',
            'date': timezone.localdate().isoformat(), 'reference': '',
        })
        p = Paiement.objects.get()
        self.assertRedirects(r, reverse('ecole:recu', args=[p.pk]))
        r = self.client.get(reverse('ecole:recu', args=[p.pk]))
        self.assertContains(r, p.numero_recu)
        self.assertContains(r, 'Vingt-cinq mille francs CFA')
        self.assertContains(r, '35 000 FCFA')  # reste à payer

    def test_encaisser_refuse_depassement_et_date_future(self):
        e = creer_eleve(montant_scolarite=10000)
        demain = (timezone.localdate() + timedelta(days=1)).isoformat()
        r = self.client.post(reverse('ecole:encaisser', args=[e.pk]), {
            'motif': 'scolarite', 'montant': '20000', 'mode': 'especes', 'date': demain,
        })
        self.assertIn('montant', r.context['form'].errors)
        self.assertIn('date', r.context['form'].errors)
        self.assertFalse(Paiement.objects.exists())

    def test_modifier_scolarite_inferieure_au_verse(self):
        e = creer_eleve()
        payer(e, 50000)
        data = {f: getattr(e, f) for f in ('nom', 'prenom', 'classe', 'parent', 'telephone')}
        data['etablissement_nom'] = e.etablissement.nom
        r = self.client.post(reverse('ecole:eleve_modifier', args=[e.pk]), {**data, 'montant_scolarite': '40000'})
        self.assertIn('montant_scolarite', r.context['form'].errors)

    def test_liste_par_classe_filtre(self):
        creer_eleve(nom='Alpha', classe='CM1')
        creer_eleve(nom='Beta', classe='CM2')
        r = self.client.get(reverse('ecole:eleves'), {'classe': 'CM1'})
        self.assertContains(r, 'ALPHA')
        self.assertNotContains(r, 'BETA')
        r = self.client.get(reverse('ecole:eleves'), {'q': 'bet'})
        self.assertContains(r, 'BETA')
        self.assertNotContains(r, 'ALPHA')

    def test_suivi_payes_impayes(self):
        payer(creer_eleve(nom='Paye'), 60000)
        payer(creer_eleve(nom='Doit'), 100)
        r = self.client.get(reverse('ecole:suivi'), {'statut': 'payes'})
        self.assertContains(r, 'PAYE')
        self.assertNotContains(r, 'DOIT')
        r = self.client.get(reverse('ecole:suivi'), {'statut': 'impayes'})
        self.assertContains(r, 'DOIT')
        self.assertContains(r, '59 900 FCFA')

    def test_journal_filtres(self):
        e = creer_eleve()
        payer(e, 1000, mode='especes')
        payer(e, 2000, mode='mobile_money', motif=Motif.TRANCHE2)
        r = self.client.get(reverse('ecole:journal'), {'mode': 'mobile_money'})
        self.assertEqual(r.context['nb'], 1)
        self.assertEqual(r.context['total'], 2000)

    def test_relance_whatsapp(self):
        e = creer_eleve(etablissement='CEG Test', telephone='97 12 34 56')
        payer(e, 10000)
        r = self.client.post(reverse('ecole:relancer', args=[e.pk, 'whatsapp']))
        self.assertEqual(r.status_code, 302)
        self.assertTrue(r['Location'].startswith('https://wa.me/22997123456?text='))
        texte = unquote(r['Location'])
        self.assertIn('50 000 FCFA', texte)
        self.assertIn('CEG Test', texte)
        self.assertEqual(Relance.objects.get().reste_a_payer, 50000)

    def test_relance_sms_et_get_refuse(self):
        e = creer_eleve()
        self.assertEqual(self.client.get(reverse('ecole:relancer', args=[e.pk, 'sms'])).status_code, 405)
        r = self.client.post(reverse('ecole:relancer', args=[e.pk, 'sms']))
        self.assertTrue(r['Location'].startswith('sms:+229'))

    def test_export_excel(self):
        e = creer_eleve(nom='Export')
        payer(e, 5000)
        r = self.client.get(reverse('ecole:export_eleves'), {'statut': 'impayes'})
        self.assertEqual(r['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws['C4'].value, 'EXPORT')
        self.assertEqual(ws['L4'].value, 55000)
        r = self.client.get(reverse('ecole:export_paiements'))
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws['G4'].value, 5000)

    def test_exports_pdf(self):
        e = creer_eleve(nom='Pdf')
        p = payer(e, 5000)
        for url, params in [(reverse('ecole:export_eleves_pdf'), {'statut': 'impayes'}),
                            (reverse('ecole:export_eleves_pdf'), {'classe': 'CM2'}),
                            (reverse('ecole:export_paiements_pdf'), {'du': '2020-01-01', 'mode': 'especes'}),
                            (reverse('ecole:recu_pdf', args=[p.pk]), {})]:
            r = self.client.get(url, params)
            self.assertEqual(r['Content-Type'], 'application/pdf', url)
            self.assertTrue(r.content.startswith(b'%PDF'), url)
        # Liste vide : le PDF est quand même produit
        r = self.client.get(reverse('ecole:export_eleves_pdf'), {'classe': 'Inexistante'})
        self.assertTrue(r.content.startswith(b'%PDF'))

    def test_pdf_interdit_au_caissier(self):
        self.client.force_login(utilisateur(roles.CAISSIER))
        self.assertEqual(self.client.get(reverse('ecole:export_eleves_pdf')).status_code, 403)

    def test_etablissement_cree_ou_reutilise(self):
        base = {'nom': 'A', 'prenom': 'B', 'classe': 'CI', 'parent': 'C', 'telephone': '1',
                'montant_scolarite': '1000', 'mode': 'especes'}
        self.client.post(reverse('ecole:eleve_nouveau'), {**base, 'etablissement_nom': '  EPP  Cadjèhoun '})
        self.client.post(reverse('ecole:eleve_nouveau'), {**base, 'etablissement_nom': 'epp cadjèhoun'})
        self.assertEqual(Etablissement.objects.get().nom, 'EPP Cadjèhoun')
        self.assertEqual(Etablissement.objects.get().eleves.count(), 2)

    def test_filtre_par_etablissement(self):
        a = creer_eleve(etablissement='École A', nom='Alpha')
        creer_eleve(etablissement='École B', nom='Beta')
        r = self.client.get(reverse('ecole:eleves'), {'etab': a.etablissement_id})
        self.assertContains(r, 'ALPHA')
        self.assertNotContains(r, 'BETA')
        self.assertContains(r, 'Tous les établissements')  # filtre visible dès 2 établissements
        r = self.client.get(reverse('ecole:suivi'), {'statut': 'impayes', 'etab': a.etablissement_id})
        self.assertEqual([e.nom for e in r.context['eleves']], ['Alpha'])
        r = self.client.get(reverse('ecole:export_eleves'), {'etab': a.etablissement_id})
        ws = load_workbook(BytesIO(r.content)).active
        self.assertEqual((ws['A4'].value, ws['C4'].value, ws['A5'].value), ('École A', 'ALPHA', None))

    def test_recu_affiche_etablissement_de_l_eleve(self):
        p = payer(creer_eleve(etablissement='Collège Les Palmiers'), 1000)
        self.assertContains(self.client.get(reverse('ecole:recu', args=[p.pk])), 'Collège Les Palmiers')

    def test_renommer_etablissement(self):
        e = creer_eleve(etablissement='Mon établissement')
        creer_eleve(etablissement='CEG Existant')
        url = reverse('ecole:etablissement_renommer', args=[e.etablissement_id])
        self.client.post(url, {'nom': 'ceg existant'})  # doublon refusé
        self.client.post(url, {'nom': ' Complexe  Scolaire Les Lauriers '})
        e.refresh_from_db()
        self.assertEqual(e.etablissement.nom, 'Complexe Scolaire Les Lauriers')
        self.client.force_login(utilisateur(roles.SECRETARIAT))
        self.assertEqual(self.client.post(url, {'nom': 'X'}).status_code, 403)

    def test_annuler_paiement(self):
        e = creer_eleve()
        p = payer(e, 1000)
        self.client.get(reverse('ecole:paiement_supprimer', args=[p.pk]))
        self.assertTrue(Paiement.objects.exists())  # un GET ne supprime rien
        self.client.post(reverse('ecole:paiement_supprimer', args=[p.pk]))
        self.assertTrue(Paiement.objects.exists())  # le motif est obligatoire
        self.client.post(reverse('ecole:paiement_supprimer', args=[p.pk]), {'motif': 'Erreur de saisie'})
        self.assertFalse(Paiement.objects.exists())

    def test_archiver_eleve_sans_rien_supprimer(self):
        e = creer_eleve()
        payer(e, 1000)
        self.client.post(reverse('ecole:eleve_archiver', args=[e.pk]))
        e.refresh_from_db()
        self.assertFalse(e.archive)  # le motif est obligatoire
        self.client.post(reverse('ecole:eleve_archiver', args=[e.pk]), {'motif': 'Transfert'})
        e.refresh_from_db()
        self.assertTrue(e.archive)
        fiche = reverse('ecole:eleve_detail', args=[e.pk])
        self.assertEqual(e.motif_archive, 'Transfert')
        self.assertTrue(Paiement.objects.filter(eleve=e).exists())  # historique conservé
        # Disparaît des listes, du suivi et de la recherche d'encaissement…
        self.assertNotContains(self.client.get(reverse('ecole:eleves')), fiche)
        self.assertNotContains(self.client.get(reverse('ecole:suivi')), fiche)
        # …mais reste consultable dans les archives et sur sa fiche
        self.assertContains(self.client.get(reverse('ecole:eleves_archives')), fiche)
        self.assertContains(self.client.get(reverse('ecole:eleve_detail', args=[e.pk])), 'Élève archivé')
        # Pas d'encaissement sur un élève archivé
        r = self.client.post(reverse('ecole:encaisser', args=[e.pk]), {'montant': 500})
        self.assertRedirects(r, reverse('ecole:eleve_detail', args=[e.pk]))
        self.assertEqual(Paiement.objects.filter(eleve=e).count(), 1)
        # Réactivation
        self.client.post(reverse('ecole:eleve_reactiver', args=[e.pk]))
        e.refresh_from_db()
        self.assertFalse(e.archive)
        self.assertContains(self.client.get(reverse('ecole:eleves')), fiche)

    def test_aucune_suppression_eleve_possible(self):
        from django.contrib.admin.sites import site
        from django.urls import NoReverseMatch
        with self.assertRaises(NoReverseMatch):
            reverse('ecole:eleve_supprimer', args=[1])
        self.assertFalse(site._registry[Eleve].has_delete_permission(None))
        for groupe in Group.objects.all():
            self.assertFalse(groupe.permissions.filter(codename='delete_eleve').exists())

    def test_creer_utilisateur_avec_role(self):
        r = self.client.post(reverse('ecole:utilisateur_nouveau'), {
            'username': 'awa', 'first_name': 'Awa', 'last_name': 'Sossa', 'role': roles.CAISSIER,
            'password1': 'Tres-Secret-2026', 'password2': 'Tres-Secret-2026',
        })
        awa = User.objects.get(username='awa')
        # On enchaîne sur sa fiche pour lui attribuer ses codes d'accès aux modules.
        self.assertRedirects(r, reverse('ecole:utilisateur_modifier', args=[awa.pk]) + '#codes')
        self.assertEqual(roles.role_de(awa), roles.CAISSIER)

    def test_direction_ne_peut_pas_se_retrograder(self):
        self.client.post(reverse('ecole:utilisateur_modifier', args=[self.user.pk]), {
            'first_name': 'X', 'last_name': 'Y', 'role': roles.CAISSIER, 'is_active': '',
        })
        self.user.refresh_from_db()
        self.assertTrue(self.user.is_active)
        self.assertEqual(roles.role_de(self.user), roles.DIRECTION)

    def test_pages_principales(self):
        e = creer_eleve()
        p = payer(e, 1000)
        payer(creer_eleve(classe='CE1'), 60000)
        urls = ['ecole:portail', 'ecole:accueil', 'ecole:eleves', 'ecole:eleve_nouveau', 'ecole:encaisser_recherche',
                'ecole:journal', 'ecole:suivi', 'ecole:relances', 'ecole:exports',
                'ecole:utilisateurs', 'ecole:utilisateur_nouveau', 'ecole:parametres']
        for name in urls:
            self.assertEqual(self.client.get(reverse(name)).status_code, 200, name)
        for name, pk in [('ecole:eleve_detail', e.pk), ('ecole:encaisser', e.pk), ('ecole:recu', p.pk),
                         ('ecole:eleve_modifier', e.pk), ('ecole:utilisateur_modifier', self.user.pk)]:
            self.assertEqual(self.client.get(reverse(name, args=[pk])).status_code, 200, name)
        self.assertEqual(self.client.get(reverse('ecole:encaisser_recherche'), {'q': 'dos'}).status_code, 200)


@override_settings(MODULES_BIENTOT=[])
class VerrouModulesTests(TestCase):
    def setUp(self):
        from . import verrou

        self.verrou = verrou
        self.user = utilisateur(roles.DIRECTION)
        self.client.force_login(self.user)

    def entrer(self, module, code):
        return self.client.post(reverse('ecole:deverrouiller', args=[module]), {'code': code, 'next': '/eleves/'})

    def test_code_toujours_demande(self):
        self.assertRedirects(self.client.get(reverse('ecole:eleves')),
                             reverse('ecole:deverrouiller', args=['scolarite']) + '?next=%2Feleves%2F')

    def test_direction_cree_son_code_a_la_premiere_entree(self):
        url = reverse('ecole:deverrouiller', args=['scolarite'])
        self.assertContains(self.client.get(url), 'Confirmez le code')
        r = self.client.post(url, {'code': '2468', 'confirmation': '2469', 'next': '/eleves/'})
        self.assertContains(r, 'pas identiques')
        r = self.client.post(url, {'code': '2468', 'confirmation': '2468', 'next': '/eleves/'})
        self.assertRedirects(r, '/eleves/')
        self.assertTrue(self.verrou.est_protege(self.user, 'scolarite'))

    def test_sans_code_pas_d_entree(self):
        self.client.force_login(utilisateur(roles.CAISSIER, 'caisse'))
        self.assertEqual(self.client.get(reverse('ecole:eleves')).status_code, 302)
        r = self.client.post(reverse('ecole:deverrouiller', args=['scolarite']),
                             {'code': '2468', 'confirmation': '2468', 'next': '/eleves/'})
        self.assertContains(r, 'Demandez-le à la direction')
        self.assertEqual(self.client.get(reverse('ecole:eleves')).status_code, 302)

    def test_code_demande_puis_accepte(self):
        self.verrou.definir_code(self.user, 'scolarite', '2468')
        r = self.client.get(reverse('ecole:eleves'))
        self.assertRedirects(r, reverse('ecole:deverrouiller', args=['scolarite']) + '?next=%2Feleves%2F')
        self.assertContains(self.entrer('scolarite', '0000'), 'Code incorrect')
        self.assertRedirects(self.entrer('scolarite', '2468'), '/eleves/')
        self.assertEqual(self.client.get(reverse('ecole:journal')).status_code, 200)

    def test_accueil_et_autre_module_referment(self):
        self.verrou.definir_code(self.user, 'scolarite', '2468')
        self.entrer('scolarite', '2468')
        self.client.get(reverse('notes:tableau'))  # autre module : la scolarité est refermée
        self.assertEqual(self.client.get(reverse('ecole:eleves')).status_code, 302)
        self.entrer('scolarite', '2468')
        self.client.get(reverse('ecole:portail'))
        self.assertEqual(self.client.get(reverse('ecole:eleves')).status_code, 302)

    def test_chaque_module_a_son_code(self):
        self.verrou.definir_code(self.user, 'scolarite', '2468')
        self.verrou.definir_code(self.user, 'notes', '1357')
        self.entrer('scolarite', '2468')
        self.assertEqual(self.client.get(reverse('notes:tableau')).status_code, 302)
        self.assertEqual(self.entrer('notes', '2468').status_code, 200)

    def test_blocage_apres_cinq_essais(self):
        self.verrou.definir_code(self.user, 'scolarite', '2468')
        for _i in range(5):
            self.entrer('scolarite', '0000')
        r = self.entrer('scolarite', '2468')
        self.assertContains(r, "Trop d'essais")
        self.assertEqual(self.client.get(reverse('ecole:eleves')).status_code, 302)

    def test_inactivite(self):
        self.verrou.definir_code(self.user, 'scolarite', '2468')
        self.entrer('scolarite', '2468')
        with self.settings(VERROU_MODULE_MINUTES=0):
            self.assertEqual(self.client.get(reverse('ecole:eleves')).status_code, 302)

    def test_redirection_externe_refusee(self):
        self.verrou.definir_code(self.user, 'scolarite', '2468')
        r = self.client.post(reverse('ecole:deverrouiller', args=['scolarite']),
                             {'code': '2468', 'next': 'https://exemple.com/'})
        self.assertRedirects(r, reverse('ecole:accueil'), fetch_redirect_response=False)

    def test_module_interdit_au_role(self):
        self.client.force_login(utilisateur(roles.CAISSIER, 'caisse'))
        self.assertEqual(self.client.get(reverse('ecole:deverrouiller', args=['admin'])).status_code, 403)

    @override_settings(VERROU_MODULES=False)  # page d'Administration : on teste le formulaire, pas le verrou
    def test_fiche_utilisateur_definit_et_retire_ses_codes(self):
        from .models import CodeModule

        caissier = utilisateur(roles.CAISSIER, 'caisse')
        url = reverse('ecole:utilisateur_modifier', args=[caissier.pk])
        self.client.post(url, {'codes': '1', 'codes-code_scolarite': '9876'})
        obj = CodeModule.objects.get(utilisateur=caissier, module='scolarite')
        self.assertNotEqual(obj.code, '9876')  # stocké chiffré
        self.assertNotContains(self.client.get(url), 'codes-code_admin')  # seulement les modules de son rôle
        self.client.post(url, {'codes': '1', 'codes-retirer_scolarite': 'on'})
        self.assertFalse(CodeModule.objects.exists())
        r = self.client.post(url, {'codes': '1', 'codes-code_scolarite': '12'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(CodeModule.objects.exists())

    def test_le_code_est_personnel(self):
        autre = utilisateur(roles.SECRETARIAT, 'secretaire')
        self.verrou.definir_code(autre, 'scolarite', '1111')
        self.verrou.definir_code(self.user, 'scolarite', '2468')
        self.assertContains(self.entrer('scolarite', '1111'), 'Code incorrect')
        self.client.force_login(autre)
        self.assertRedirects(self.entrer('scolarite', '1111'), '/eleves/')


class ConnexionPuisModuleTests(TestCase):
    def test_connexion_simple_puis_code_du_module(self):
        from . import verrou

        user = utilisateur(roles.SECRETARIAT, 'secretaire')
        verrou.definir_code(user, 'scolarite', '2468')
        r = self.client.post(reverse('login'), {'username': 'secretaire', 'password': 'Mot-de-passe-123'})
        self.assertRedirects(r, reverse('ecole:portail'))  # identifiant + mot de passe suffisent
        r = self.client.get(reverse('ecole:accueil'))  # clic sur la tuile Scolarité
        self.assertRedirects(r, reverse('ecole:deverrouiller', args=['scolarite']) + '?next=%2Fscolarite%2F')
        r = self.client.post(reverse('ecole:deverrouiller', args=['scolarite']), {'code': '2468', 'next': '/scolarite/'})
        self.assertRedirects(r, reverse('ecole:accueil'))


@override_settings(MODULES_BIENTOT=['notes', 'profs', 'examens'])
class ModulesBientotTests(TestCase):
    def setUp(self):
        self.user = utilisateur(roles.DIRECTION)
        self.client.force_login(self.user)

    def test_tuiles_grisees_sur_l_accueil(self):
        r = self.client.get(reverse('ecole:portail'))
        self.assertContains(r, 'Bientôt disponible', count=3)
        self.assertNotContains(r, f'href="{reverse("notes:tableau")}"')
        self.assertContains(r, f'href="{reverse("ecole:accueil")}"')

    def test_pages_fermees(self):
        for url in (reverse('notes:tableau'), reverse('profs:liste'), reverse('examens:liste'),
                    reverse('ecole:deverrouiller', args=['notes'])):
            self.assertContains(self.client.get(url), 'Ce module sera ouvert prochainement', msg_prefix=url)

    @override_settings(VERROU_MODULES=False)
    def test_pas_de_code_pour_un_module_ferme(self):
        r = self.client.get(reverse('ecole:utilisateur_modifier', args=[self.user.pk]))
        self.assertNotContains(r, 'codes-code_notes')

    def test_scolarite_reste_ouverte(self):
        from . import verrou

        verrou.definir_code(self.user, 'scolarite', '2468')
        r = self.client.post(reverse('ecole:deverrouiller', args=['scolarite']), {'code': '2468', 'next': '/eleves/'})
        self.assertRedirects(r, '/eleves/')
