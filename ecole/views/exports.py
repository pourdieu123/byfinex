from io import BytesIO

from django.conf import settings
from django.contrib.auth.decorators import login_required, permission_required
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.functional import Promise
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import pdf
from ..models import Eleve, Etablissement, ModePaiement, Motif
from .eleves import classes_existantes, etab_demande, etablissements_utilises
from .paiements import filtrer_paiements

FORMAT_MONTANT = '#,##0" FCFA"'



def _cellule_sure(valeur):
    """Empêche l'injection de formules : un texte commençant par = + - @ est rendu inerte dans Excel."""
    if isinstance(valeur, str) and valeur[:1] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + valeur
    return valeur

def _etab_filtre(request):
    etab_id = etab_demande(request)
    return Etablissement.objects.filter(pk=etab_id).first() if etab_id else None


def _classeur(titre, entetes, lignes, colonnes_montant, etab=None):
    wb = Workbook()
    ws = wb.active
    ws.title = titre[:31]
    ws.append([f'{etab.nom if etab else settings.APP_NOM} — {titre} — {timezone.localdate():%d/%m/%Y}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(entetes)
    for cell in ws[3]:
        cell.font = Font(bold=True, color='0A0A0A')
        cell.fill = PatternFill('solid', fgColor='EDEDED')
        cell.alignment = Alignment(vertical='center')
    for ligne in lignes:
        # Les libellés traduits « paresseux » (statut, mode…) doivent être convertis en texte pour Excel.
        ws.append([_cellule_sure(str(v) if isinstance(v, Promise) else v) for v in ligne])
    for idx, entete in enumerate(entetes, start=1):
        lettre = get_column_letter(idx)
        largeur = max([len(str(entete))] + [len(str(l[idx - 1])) for l in lignes]) + 2
        ws.column_dimensions[lettre].width = min(max(largeur, 10), 45)
        if idx in colonnes_montant:
            for (cell,) in ws.iter_rows(min_row=4, min_col=idx, max_col=idx):
                cell.number_format = FORMAT_MONTANT
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:{get_column_letter(len(entetes))}{3 + len(lignes)}'
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _reponse(contenu, nom):
    reponse = HttpResponse(
        contenu, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    reponse['Content-Disposition'] = f'attachment; filename="{nom}-{timezone.localdate():%Y%m%d}.xlsx"'
    return reponse


def _eleves_filtres(request, statut, etab):
    base = Eleve.objects.actifs().filter(etablissement=etab) if etab else Eleve.objects.actifs()
    qs = {'payes': base.soldes, 'impayes': base.impayes}.get(statut, base.avec_totaux)()
    if request.GET.get('classe'):
        qs = qs.filter(classe=request.GET['classe'])
    return qs.select_related('etablissement')


def _seul_etablissement(eleves):
    """L'établissement commun à tous ces élèves, s'il n'y en a qu'un (pour l'en-tête du PDF)."""
    etabs = {e.etablissement for e in eleves}
    return etabs.pop() if len(etabs) == 1 else None


@login_required
@permission_required('ecole.exporter_donnees', raise_exception=True)
def exports(request):
    return render(request, 'ecole/exports.html', {
        'classes': classes_existantes(),
        'etablissements': etablissements_utilises(),
        'modes': ModePaiement.choices,
        'motifs': Motif.choices,
    })


@login_required
@permission_required('ecole.exporter_donnees', raise_exception=True)
def export_eleves(request):
    statut = request.GET.get('statut', '')
    etab = _etab_filtre(request)
    qs = _eleves_filtres(request, statut, etab)
    titre = {'payes': _('Élèves soldés'), 'impayes': _('Élèves impayés')}.get(statut, _('Liste des élèves'))
    entetes = [_('Établissement'), _('Classe'), _('Nom'), _('Prénom(s)'), _('N° Educ Master'), _('Parent / tuteur'), _('Nature'), _('Téléphone'),
               _('Résidence'), _('Scolarité'), _('Versé (tranches)'), _('Reste à payer'), _('Inscription due'),
               _('Inscription versée'), _('Inscription'), _('Statut')]
    lignes = [
        [e.etablissement.nom, e.classe, e.nom.upper(), e.prenom, e.educ_master, e.parent, e.get_nature_parent_display(),
         e.telephone, e.residence, e.montant_scolarite, e.total_verse, e.reste_a_payer,
         e.inscription_due, e.total_inscription, e.statut_inscription_libelle, e.statut_libelle]
        for e in qs
    ]
    return _reponse(_classeur(titre, entetes, lignes, colonnes_montant={10, 11, 12, 13, 14}, etab=etab), 'eleves')


@login_required
@permission_required('ecole.exporter_donnees', raise_exception=True)
def export_paiements(request):
    entetes = [_('N° reçu'), _('Date'), _('Élève'), _('Établissement'), _('Classe'), _('Motif'), _('Montant'), _('Mode'), _('Référence'), _('Encaissé par')]
    lignes = [
        [p.numero_recu, p.date, str(p.eleve), p.eleve.etablissement.nom, p.eleve.classe, p.get_motif_display(), p.montant,
         p.get_mode_display(), p.reference,
         (p.encaisse_par.get_full_name() or p.encaisse_par.username) if p.encaisse_par else '']
        for p in filtrer_paiements(request.GET)
    ]
    contenu = _classeur(_('Journal des paiements'), entetes, lignes, colonnes_montant={7}, etab=_etab_filtre(request))
    return _reponse(contenu, 'paiements')


# --- Exports PDF ---

def _reponse_pdf(contenu, nom, telecharger=True):
    reponse = HttpResponse(contenu, content_type='application/pdf')
    mode = 'attachment' if telecharger else 'inline'
    reponse['Content-Disposition'] = f'{mode}; filename="{nom}-{timezone.localdate():%Y%m%d}.pdf"'
    return reponse


def _periode(params):
    du, au = params.get('du'), params.get('au')
    fmt = lambda d: '/'.join(reversed(d.split('-')))  # noqa: E731 — 2026-09-24 -> 24/09/2026
    if du and au:
        return _('du %(du)s au %(au)s') % {'du': fmt(du), 'au': fmt(au)}
    if du:
        return _('depuis le %(du)s') % {'du': fmt(du)}
    if au:
        return _('jusqu’au %(au)s') % {'au': fmt(au)}
    return _('toutes dates')


@login_required
@permission_required('ecole.exporter_donnees', raise_exception=True)
def export_eleves_pdf(request):
    statut = request.GET.get('statut', '')
    classe = request.GET.get('classe', '')
    etab = _etab_filtre(request)
    qs = _eleves_filtres(request, statut, etab)
    titre = {'payes': _('Élèves ayant payé'), 'impayes': _('Élèves n’ayant pas payé')}.get(statut, _('Liste des élèves'))
    sous_titre = _('%(valeur)s · classés par classe puis par nom') % {'valeur': "Classe " + classe if classe else "Toutes les classes"}
    contenu = pdf.pdf_eleves(qs, etab or _seul_etablissement(qs), titre, sous_titre)
    return _reponse_pdf(contenu, {'payes': 'eleves-payes', 'impayes': 'eleves-impayes'}.get(statut, 'eleves'))


@login_required
@permission_required('ecole.exporter_donnees', raise_exception=True)
def export_paiements_pdf(request):
    filtres = [_periode(request.GET)]
    if request.GET.get('motif') in Motif.values:
        filtres.append(Motif(request.GET['motif']).label)
    if request.GET.get('mode') in ModePaiement.values:
        filtres.append(ModePaiement(request.GET['mode']).label)
    if request.GET.get('classe'):
        filtres.append(_('Classe %(GET)s') % {'GET': request.GET["classe"]})
    paiements = list(filtrer_paiements(request.GET))
    etab = _etab_filtre(request) or _seul_etablissement([p.eleve for p in paiements])
    contenu = pdf.pdf_paiements(paiements, etab, ' · '.join(str(f) for f in filtres))
    return _reponse_pdf(contenu, 'paiements')
